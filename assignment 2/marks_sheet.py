from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

FILE_NAME = "Python_Class_Marks.xlsx"
TAB_NAME = "Marks"
CUTOFF = 40

COLUMNS = [
    "Roll No",
    "Name",
    "Practical (out of 50)",
    "Written (out of 50)",
    "Total (out of 100)",
    "Status",
]

RECORDS = [
    (201, "Isha Kulkarni", 44, 39),
    (202, "Yash Thakkar", 31, 22),
    (203, "Meera Joshi", 19, 14),
    (204, "Neil Bhatt", 46, 41),
    (205, "Tanvi Rao", 27, 10),
    (206, "Harsh Vora", 36, 38),
    (207, "Pooja Nair", 22, 25),
]


def new_book():
    book = Workbook()
    tab = book.active
    tab.title = TAB_NAME
    return book, tab


def fill_rows(tab):
    tab.append(COLUMNS)
    for roll, name, practical, written in RECORDS:
        tab.append([roll, name, practical, written])


def compute(tab):
    totals = []
    for r in range(2, tab.max_row + 1):
        practical = tab.cell(row=r, column=3).value
        written = tab.cell(row=r, column=4).value
        total = practical + written
        tab.cell(row=r, column=5, value=total)
        tab.cell(row=r, column=6, value="Pass" if total >= CUTOFF else "Fail")
        totals.append(total)
    return totals


def add_average(tab, totals):
    mean = round(sum(totals) / len(totals), 2)
    r = tab.max_row + 1
    tab.cell(row=r, column=2, value="Average")
    tab.cell(row=r, column=5, value=mean)
    tab.cell(row=r, column=2).font = Font(bold=True)
    tab.cell(row=r, column=5).font = Font(bold=True)
    return mean


def format_header(tab):
    fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for cell in tab[1]:
        cell.font = Font(bold=True)
        cell.fill = fill


def fit_columns(tab):
    for col in tab.columns:
        widest = max(len(str(c.value)) for c in col if c.value is not None)
        tab.column_dimensions[col[0].column_letter].width = widest + 3


def main():
    book, tab = new_book()
    fill_rows(tab)
    totals = compute(tab)
    mean = add_average(tab, totals)
    format_header(tab)
    fit_columns(tab)
    book.save(FILE_NAME)
    print("Created", FILE_NAME)
    print("Average total:", mean)


if __name__ == "__main__":
    main()
